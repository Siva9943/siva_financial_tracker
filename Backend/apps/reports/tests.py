import io
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.budgets.models import Budget
from apps.investments.models import Investment, InvestmentTransaction
from apps.loans.models import Loan
from apps.planning.models import PlanItem, RepaymentPlan
from apps.transactions.models import Transaction
from services.investment_service import record_investment_transaction
from services.report_export import export_to_csv, export_to_excel, export_to_pdf
from services.report_service import (
    ReportError,
    build_budget_report,
    build_dividend_report,
    build_expense_report,
    build_investment_report,
    build_investment_transaction_report,
    build_loan_report,
    build_monthly_report,
    build_portfolio_performance_report,
    build_repayment_report,
    build_yearly_report,
)

User = get_user_model()


class ReportBuilderTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='alice', password='pass12345')
        Transaction.objects.create(
            user=self.user, transaction_type=Transaction.TransactionType.INCOME,
            amount=Decimal('5000'), category='Salary', transaction_date=date(2026, 3, 5),
        )
        Transaction.objects.create(
            user=self.user, transaction_type=Transaction.TransactionType.EXPENSE,
            amount=Decimal('1200'), category='Food', transaction_date=date(2026, 3, 10),
        )
        self.loan = Loan.objects.create(
            user=self.user, loan_name='Home Loan', loan_type=Loan.LoanType.HOME,
            original_principal=Decimal('100000'), outstanding_principal=Decimal('90000'),
            annual_interest_rate=Decimal('10'), interest_calculation_method=Loan.InterestMethod.EMI_AMORTIZATION,
            emi_amount=Decimal('5000'), start_date=date(2026, 1, 1), maturity_date=date(2030, 1, 1),
        )

    def test_monthly_report_structure_and_totals(self):
        report = build_monthly_report(self.user, month=3, year=2026)
        self.assertIn('March 2026', report['title'])
        summary = report['sections'][0]
        summary_dict = dict(summary['rows'])
        self.assertEqual(summary_dict['Total Income'], Decimal('5000.00'))
        self.assertEqual(summary_dict['Total Expenses'], Decimal('1200.00'))
        self.assertEqual(summary_dict['Savings'], Decimal('3800.00'))

    def test_yearly_report_covers_full_year(self):
        report = build_yearly_report(self.user, year=2026)
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Total Income'], Decimal('5000.00'))

    def test_loan_report_lists_all_loans(self):
        report = build_loan_report(self.user)
        loan_rows = report['sections'][1]['rows']
        self.assertEqual(len(loan_rows), 1)
        self.assertEqual(loan_rows[0][0], 'Home Loan')

    def test_expense_report_breakdown_and_transactions(self):
        report = build_expense_report(self.user, date(2026, 3, 1), date(2026, 3, 31))
        breakdown_rows = report['sections'][0]['rows']
        transaction_rows = report['sections'][1]['rows']
        self.assertEqual(breakdown_rows[0][0], 'Food')
        self.assertEqual(len(transaction_rows), 1)

    def test_budget_report(self):
        Budget.objects.create(user=self.user, category='Food', amount=Decimal('1000'), month=3, year=2026)
        report = build_budget_report(self.user, month=3, year=2026)
        rows = report['sections'][0]['rows']
        self.assertEqual(rows[0][0], 'Food')
        self.assertEqual(rows[0][2], Decimal('1200.00'))  # spent

    def test_repayment_report_raises_when_no_plan(self):
        with self.assertRaises(ReportError):
            build_repayment_report(self.user)

    def test_repayment_report_with_plan(self):
        plan = RepaymentPlan.objects.create(
            user=self.user, strategy='AVALANCHE', monthly_income=Decimal('50000'), essential_expenses=Decimal('20000'),
            total_existing_emi=Decimal('5000'), months_remaining=12, expected_payoff_date=date(2027, 1, 1),
            total_interest=Decimal('6000'), total_payment=Decimal('66000'),
        )
        PlanItem.objects.create(
            plan=plan, month_number=1, month_date=date(2026, 4, 1), loan=self.loan, loan_name='Home Loan',
            normal_emi=Decimal('5000'), principal_paid=Decimal('4000'), interest_paid=Decimal('1000'), closing_balance=Decimal('86000'),
        )
        report = build_repayment_report(self.user)
        self.assertEqual(len(report['sections'][1]['rows']), 1)

    def test_reports_only_use_own_users_data(self):
        other_user = User.objects.create_user(username='bob', password='pass12345')
        Transaction.objects.create(
            user=other_user, transaction_type=Transaction.TransactionType.INCOME,
            amount=Decimal('999999'), category='Salary', transaction_date=date(2026, 3, 5),
        )
        report = build_monthly_report(self.user, month=3, year=2026)
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Total Income'], Decimal('5000.00'))


class InvestmentReportBuilderTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='alice', password='pass12345')
        self.investment = Investment.objects.create(
            user=self.user, name='Reliance Industries', investment_type=Investment.InvestmentType.STOCK, current_price=Decimal('1520'),
        )
        record_investment_transaction(
            self.investment, transaction_type=InvestmentTransaction.TransactionType.BUY,
            transaction_date=date(2026, 1, 10), quantity=Decimal('10'), price=Decimal('1400'), today=date(2026, 3, 18),
        )
        record_investment_transaction(
            self.investment, transaction_type=InvestmentTransaction.TransactionType.DIVIDEND,
            transaction_date=date(2026, 2, 1), amount=Decimal('250'), today=date(2026, 3, 18),
        )

    def test_investment_report_lists_holdings(self):
        report = build_investment_report(self.user)
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Total Invested'], Decimal('14000.00'))
        holdings_rows = report['sections'][1]['rows']
        self.assertEqual(len(holdings_rows), 1)
        self.assertEqual(holdings_rows[0][0], 'Reliance Industries')

    def test_portfolio_performance_report_breaks_out_return_components(self):
        report = build_portfolio_performance_report(self.user)
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Dividend Income'], Decimal('250.00'))
        self.assertEqual(summary_dict['Methodology'], 'AVERAGE_COST_SIMPLE_RETURN')
        allocation_rows = report['sections'][1]['rows']
        self.assertEqual(allocation_rows[0][0], 'STOCK')

    def test_dividend_report_lists_payouts_and_total(self):
        report = build_dividend_report(self.user)
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Total Dividend Income'], Decimal('250.00'))
        self.assertEqual(len(report['sections'][1]['rows']), 1)

    def test_dividend_report_respects_date_range(self):
        report = build_dividend_report(self.user, start=date(2026, 3, 1), end=date(2026, 3, 31))
        summary_dict = dict(report['sections'][0]['rows'])
        self.assertEqual(summary_dict['Total Dividend Income'], Decimal('0.00'))

    def test_investment_transaction_report_lists_all_transaction_types(self):
        report = build_investment_transaction_report(self.user, start=date(2026, 1, 1), end=date(2026, 3, 31))
        rows = report['sections'][0]['rows']
        self.assertEqual(len(rows), 2)

    def test_reports_only_include_own_investments(self):
        other_user = User.objects.create_user(username='bob', password='pass12345')
        Investment.objects.create(user=other_user, name='Other', investment_type=Investment.InvestmentType.STOCK, current_price=Decimal('100'))
        report = build_investment_report(self.user)
        self.assertEqual(len(report['sections'][1]['rows']), 1)


class ReportExportTests(TestCase):
    def setUp(self):
        self.report = {
            'title': 'Test Report',
            'generated_at': date(2026, 3, 19),
            'sections': [{'heading': 'Numbers', 'columns': ['A', 'B'], 'rows': [[1, 2], [3, 4]]}],
        }

    def test_csv_export_round_trips(self):
        content = export_to_csv(self.report)
        text = content.decode('utf-8-sig')
        self.assertIn('Test Report', text)
        self.assertIn('Numbers', text)
        self.assertIn('1,2', text)

    def test_excel_export_is_readable(self):
        content = export_to_excel(self.report)
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
        self.assertIn('Test Report', values)
        self.assertIn('Numbers', values)

    def test_pdf_export_produces_valid_pdf_bytes(self):
        content = export_to_pdf(self.report)
        self.assertTrue(content.startswith(b'%PDF'))

    def test_empty_section_does_not_crash_pdf_export(self):
        report = {'title': 'Empty', 'generated_at': date(2026, 3, 19), 'sections': [{'heading': 'Nothing', 'columns': ['A'], 'rows': []}]}
        content = export_to_pdf(report)
        self.assertTrue(content.startswith(b'%PDF'))


class ReportAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='alice', password='pass12345')
        self.other_user = User.objects.create_user(username='bob', password='pass12345')
        self.client.force_authenticate(self.user)
        Transaction.objects.create(
            user=self.user, transaction_type=Transaction.TransactionType.INCOME,
            amount=Decimal('5000'), category='Salary', transaction_date=date(2026, 3, 5),
        )

    def test_unknown_report_type_rejected(self):
        response = self.client.get(reverse('report', args=['unknown']))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_investment_report_json(self):
        response = self.client.get(reverse('report', args=['investment']))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('report', response.data)

    def test_portfolio_performance_report_json(self):
        response = self.client.get(reverse('report', args=['portfolio_performance']))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_dividend_report_json(self):
        response = self.client.get(reverse('report', args=['dividend']))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_investment_transaction_report_json(self):
        response = self.client.get(reverse('report', args=['investment_transaction']))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_investment_report_pdf_download(self):
        response = self.client.get(reverse('report', args=['investment']), {'export': 'pdf'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_monthly_report_json(self):
        response = self.client.get(reverse('report', args=['monthly']), {'month': 3, 'year': 2026})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('report', response.data)

    def test_monthly_report_csv_download(self):
        response = self.client.get(reverse('report', args=['monthly']), {'month': 3, 'year': 2026, 'export': 'csv'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertIn('attachment', response['Content-Disposition'])

    def test_monthly_report_xlsx_download(self):
        response = self.client.get(reverse('report', args=['monthly']), {'month': 3, 'year': 2026, 'export': 'xlsx'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', response['Content-Type'])

    def test_monthly_report_pdf_download(self):
        response = self.client.get(reverse('report', args=['monthly']), {'month': 3, 'year': 2026, 'export': 'pdf'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_unknown_format_rejected(self):
        response = self.client.get(reverse('report', args=['monthly']), {'export': 'docx'})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_repayment_report_without_plan_returns_400(self):
        response = self.client.get(reverse('report', args=['repayment']))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Repayment Planner', response.data['message'])

    def test_expense_report_invalid_range_rejected(self):
        response = self.client.get(reverse('report', args=['expense']), {'start': '2026-03-31', 'end': '2026-03-01'})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_report_isolates_users(self):
        Transaction.objects.create(
            user=self.other_user, transaction_type=Transaction.TransactionType.INCOME,
            amount=Decimal('999999'), category='Salary', transaction_date=date(2026, 3, 5),
        )
        response = self.client.get(reverse('report', args=['monthly']), {'month': 3, 'year': 2026})
        summary_rows = response.data['report']['sections'][0]['rows']
        summary_dict = dict(summary_rows)
        self.assertEqual(summary_dict['Total Income'], Decimal('5000.00'))

    def test_unauthenticated_rejected(self):
        self.client.force_authenticate(None)
        response = self.client.get(reverse('report', args=['monthly']))
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
