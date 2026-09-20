"""Generic renderers for the {title, sections} report shape (skill §47).

One implementation per format, reused by every report type — a report
builder never needs to know how it will be exported.
"""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _stringify_row(row):
    return [str(cell) for cell in row]


def export_to_csv(report):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([report['title']])
    writer.writerow([f'Generated {report["generated_at"]}'])
    writer.writerow([])

    for section in report['sections']:
        writer.writerow([section['heading']])
        writer.writerow(section['columns'])
        for row in section['rows']:
            writer.writerow(_stringify_row(row))
        writer.writerow([])

    return buffer.getvalue().encode('utf-8-sig')  # BOM so Excel opens UTF-8 CSVs correctly


def export_to_excel(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report['title'][:31] or 'Report'

    bold = Font(bold=True)
    sheet.append([report['title']])
    sheet['A1'].font = Font(bold=True, size=14)
    sheet.append([f'Generated {report["generated_at"]}'])
    sheet.append([])

    for section in report['sections']:
        heading_row = sheet.max_row + 1
        sheet.append([section['heading']])
        sheet.cell(row=heading_row, column=1).font = bold

        header_row = sheet.max_row + 1
        sheet.append(section['columns'])
        for cell in sheet[header_row]:
            cell.font = bold

        for row in section['rows']:
            sheet.append(_stringify_row(row))
        sheet.append([])

    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_to_pdf(report):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph(report['title'], styles['Title']), Paragraph(f'Generated {report["generated_at"]}', styles['Normal']), Spacer(1, 12)]

    for section in report['sections']:
        elements.append(Paragraph(section['heading'], styles['Heading2']))
        table_data = [section['columns']] + [_stringify_row(row) for row in section['rows']]
        if len(table_data) == 1:
            table_data.append(['—'] * len(section['columns']))

        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1c47d6')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTSIZE', (0, 0), (-1, -1), 8),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e1e0d9')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9f9f7')]),
                ]
            )
        )
        elements.append(table)
        elements.append(Spacer(1, 16))

    doc.build(elements)
    return buffer.getvalue()


EXPORTERS = {
    'csv': (export_to_csv, 'text/csv', 'csv'),
    'xlsx': (export_to_excel, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'),
    'pdf': (export_to_pdf, 'application/pdf', 'pdf'),
}
